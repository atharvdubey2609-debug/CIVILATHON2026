from flask import Flask, render_template, request, redirect, url_for, flash, send_from_directory, jsonify, session, abort
from werkzeug.utils import secure_filename
from pathlib import Path
from datetime import datetime
from zoneinfo import ZoneInfo
import sqlite3
import csv
import io
import uuid
import os

BASE_DIR = Path(__file__).resolve().parent
# Railway volume is mounted at /data. Use it when available so registrations
# and uploaded files survive deployments; keep BASE_DIR for local development.
DATA_DIR = Path(os.getenv('CIVILATHON_DATA_DIR', '/data' if Path('/data').is_dir() else str(BASE_DIR))).resolve()
DATA_DIR.mkdir(parents=True, exist_ok=True)

DB_PATH = DATA_DIR / 'civilathon.db'
UPLOAD_DIR = DATA_DIR / 'uploads' / 'ppt'
PAYMENT_PROOF_DIR = DATA_DIR / 'uploads' / 'payment_proofs'

# Best-effort migration from the old application directory if files still exist.
# Never overwrite data already present on the persistent volume.
import shutil
_old_db = BASE_DIR / 'civilathon.db'
if DB_PATH != _old_db and not DB_PATH.exists() and _old_db.is_file():
    shutil.copy2(_old_db, DB_PATH)
for _old_dir, _new_dir in [
    (BASE_DIR / 'uploads' / 'ppt', UPLOAD_DIR),
    (BASE_DIR / 'uploads' / 'payment_proofs', PAYMENT_PROOF_DIR),
]:
    if _old_dir.resolve() != _new_dir.resolve() and _old_dir.is_dir() and not _new_dir.exists():
        _new_dir.parent.mkdir(parents=True, exist_ok=True)
        shutil.copytree(_old_dir, _new_dir)

UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
PAYMENT_PROOF_DIR.mkdir(parents=True, exist_ok=True)

ALLOWED_EXTENSIONS = {'.ppt', '.pptx'}
ALLOWED_PROOF_EXTENSIONS = {'.png', '.jpg', '.jpeg', '.webp', '.pdf'}
MAX_FILE_SIZE = 25 * 1024 * 1024
MAX_PROOF_SIZE = 8 * 1024 * 1024
TEAM_FEE = 15000      # ₹150 in paise
INDIVIDUAL_FEE = 5000 # ₹50 in paise

app = Flask(__name__)
app.secret_key = os.getenv('FLASK_SECRET_KEY', 'local-development-only-change-this-secret')
app.config['MAX_CONTENT_LENGTH'] = MAX_FILE_SIZE + MAX_PROOF_SIZE

DOMAINS = [
    'Smart Construction & Construction Management',
    'Sustainable & Green Construction',
    'Innovative Construction Materials',
    'AI/ML & Automation',
    'IoT & Smart Infrastructure',
    'GIS, Surveying & Geospatial Technology',
    'Transportation & Smart Mobility',
    'Water Resources & Management',
    'Smart Cities & Urban Infrastructure',
    'Structural Engineering & Safety',
    'Waste Management & Circular Construction',
    'Any Other Real-World Civil Engineering Problem',
]

UPI_ID = os.getenv('CIVILATHON_UPI_ID', '9713559328@axl').strip()
ADMIN_PASSWORD = os.getenv('CIVILATHON_ADMIN_PASSWORD', '').strip()
PAYMENT_QR_PATH = BASE_DIR / 'static' / 'images' / 'payment-qr.png'


def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    conn = get_db()
    conn.execute('''
        CREATE TABLE IF NOT EXISTS registrations (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            registration_id TEXT UNIQUE NOT NULL,
            participation_type TEXT NOT NULL,
            team_name TEXT,
            college TEXT NOT NULL,
            leader_name TEXT NOT NULL,
            leader_enrollment TEXT NOT NULL DEFAULT '',
            leader_mobile TEXT NOT NULL,
            leader_email TEXT NOT NULL,
            member_2 TEXT,
            member_2_enrollment TEXT,
            member_3 TEXT,
            member_3_enrollment TEXT,
            member_4 TEXT,
            member_4_enrollment TEXT,
            research_domain TEXT NOT NULL,
            problem_statement TEXT NOT NULL,
            proposed_solution TEXT NOT NULL,
            innovation TEXT NOT NULL,
            expected_impact TEXT NOT NULL,
            ppt_filename TEXT NOT NULL,
            ppt_path TEXT NOT NULL,
            payment_status TEXT NOT NULL DEFAULT 'Pending',
            payment_id TEXT,
            payment_order_id TEXT,
            payment_signature TEXT,
            amount_paise INTEGER,
            created_at TEXT NOT NULL,
            paid_at TEXT,
            payment_reference TEXT,
            payment_screenshot_path TEXT,
            payment_submitted_at TEXT,
            admin_note TEXT
        )
    ''')
    # Safe migration for databases created by earlier Stage 3 builds.
    columns = {row['name'] for row in conn.execute('PRAGMA table_info(registrations)').fetchall()}
    migrations = {
        'payment_order_id': 'ALTER TABLE registrations ADD COLUMN payment_order_id TEXT',
        'payment_signature': 'ALTER TABLE registrations ADD COLUMN payment_signature TEXT',
        'amount_paise': 'ALTER TABLE registrations ADD COLUMN amount_paise INTEGER',
        'paid_at': 'ALTER TABLE registrations ADD COLUMN paid_at TEXT',
        'payment_reference': 'ALTER TABLE registrations ADD COLUMN payment_reference TEXT',
        'payment_screenshot_path': 'ALTER TABLE registrations ADD COLUMN payment_screenshot_path TEXT',
        'payment_submitted_at': 'ALTER TABLE registrations ADD COLUMN payment_submitted_at TEXT',
        'admin_note': 'ALTER TABLE registrations ADD COLUMN admin_note TEXT',
        'leader_enrollment': "ALTER TABLE registrations ADD COLUMN leader_enrollment TEXT NOT NULL DEFAULT ''",
        'member_2_enrollment': 'ALTER TABLE registrations ADD COLUMN member_2_enrollment TEXT',
        'member_3_enrollment': 'ALTER TABLE registrations ADD COLUMN member_3_enrollment TEXT',
        'member_4_enrollment': 'ALTER TABLE registrations ADD COLUMN member_4_enrollment TEXT',
    }
    for name, sql in migrations.items():
        if name not in columns:
            conn.execute(sql)
    conn.commit()
    conn.close()


def allowed_file(filename):
    return Path(filename).suffix.lower() in ALLOWED_EXTENSIONS


def make_registration_id():
    return f"CIV-26-{uuid.uuid4().hex[:6].upper()}"


def allowed_proof(filename):
    return Path(filename).suffix.lower() in ALLOWED_PROOF_EXTENSIONS


@app.context_processor
def inject_common():
    return {'event_domains': DOMAINS, 'upi_id': UPI_ID}


@app.route('/')
def home():
    return render_template('index.html')


@app.get('/about')
def about():
    return render_template('about.html')

@app.get('/domains')
def domains():
    return render_template('domains.html')

@app.get('/rules')
def rules():
    return render_template('rules.html')

@app.get('/guidelines')
def guidelines():
    return render_template('guidelines.html')

@app.get('/submission')
def submission():
    return render_template('submission.html')


@app.route('/register', methods=['GET', 'POST'])
def register():
    if request.method == 'GET':
        return render_template('register.html', domains=DOMAINS)

    participation_type = request.form.get('participation_type', '').strip().lower()
    team_name = request.form.get('team_name', '').strip()
    college = request.form.get('college', '').strip()
    leader_name = request.form.get('leader_name', '').strip()
    leader_enrollment = request.form.get('leader_enrollment', '').strip()
    leader_mobile = request.form.get('leader_mobile', '').strip()
    leader_email = request.form.get('leader_email', '').strip()
    member_2 = request.form.get('member_2', '').strip()
    member_2_enrollment = request.form.get('member_2_enrollment', '').strip()
    member_3 = request.form.get('member_3', '').strip()
    member_3_enrollment = request.form.get('member_3_enrollment', '').strip()
    member_4 = request.form.get('member_4', '').strip()
    member_4_enrollment = request.form.get('member_4_enrollment', '').strip()
    research_domain = request.form.get('research_domain', '').strip()
    problem_statement = request.form.get('problem_statement', '').strip()
    proposed_solution = request.form.get('proposed_solution', '').strip()
    innovation = request.form.get('innovation', '').strip()
    expected_impact = request.form.get('expected_impact', '').strip()
    ppt = request.files.get('ppt')

    errors = []
    if participation_type not in {'team', 'individual'}:
        errors.append('Please select team or individual participation.')
    if not college:
        errors.append('College name is required.')
    if not leader_name:
        errors.append('Participant/team leader name is required.')
    if not leader_enrollment:
        errors.append('Team leader/participant enrollment number is required.')
    if not leader_mobile or not leader_mobile.replace('+', '').replace(' ', '').isdigit():
        errors.append('Enter a valid mobile number.')
    if not leader_email or '@' not in leader_email:
        errors.append('Enter a valid email address.')
    if not research_domain or research_domain not in DOMAINS:
        errors.append('Please select a valid research domain.')
    for label, value in [
        ('problem statement', problem_statement),
        ('proposed solution', proposed_solution),
        ('innovation', innovation),
        ('expected impact', expected_impact),
    ]:
        if not value:
            errors.append(f'{label.capitalize()} is required.')

    if participation_type == 'team':
        if not team_name:
            errors.append('Team name is required for team participation.')
        members = [
            (member_2, member_2_enrollment, 'Member 2'),
            (member_3, member_3_enrollment, 'Member 3'),
            (member_4, member_4_enrollment, 'Member 4'),
        ]
        if not any(name for name, _enrollment, _label in members):
            errors.append('Add at least one additional team member for team participation.')
        for name, enrollment, label in members:
            if name and not enrollment:
                errors.append(f'{label} enrollment number is required when their name is entered.')
            if enrollment and not name:
                errors.append(f'Enter the name for {label} with the supplied enrollment number.')
    else:
        team_name = ''
        member_2 = member_3 = member_4 = ''
        member_2_enrollment = member_3_enrollment = member_4_enrollment = ''

    if ppt and ppt.filename and not allowed_file(ppt.filename):
        errors.append('Only .ppt and .pptx files are accepted.')

    if errors:
        for error in errors:
            flash(error, 'error')
        return render_template('register.html', domains=DOMAINS), 400

    registration_id = make_registration_id()
    original_name = ''
    stored_path = ''
    if ppt and ppt.filename:
        original_name = secure_filename(ppt.filename)
        stored_name = f'{registration_id}_{original_name}'
        file_path = UPLOAD_DIR / stored_name
        ppt.save(file_path)
        stored_path = str(file_path)

    amount_paise = TEAM_FEE if participation_type == 'team' else INDIVIDUAL_FEE
    created_at = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    conn = get_db()
    conn.execute('''
        INSERT INTO registrations (
            registration_id, participation_type, team_name, college,
            leader_name, leader_enrollment, leader_mobile, leader_email,
            member_2, member_2_enrollment, member_3, member_3_enrollment,
            member_4, member_4_enrollment, research_domain, problem_statement,
            proposed_solution, innovation, expected_impact, ppt_filename, ppt_path,
            payment_status, amount_paise, created_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    ''', (
        registration_id, participation_type, team_name, college,
        leader_name, leader_enrollment, leader_mobile, leader_email,
        member_2, member_2_enrollment, member_3, member_3_enrollment,
        member_4, member_4_enrollment, research_domain, problem_statement,
        proposed_solution, innovation, expected_impact, original_name, str(stored_path),
        'Pending', amount_paise, created_at
    ))
    conn.commit()
    conn.close()

    return redirect(url_for('payment', registration_id=registration_id))


@app.route('/payment/<registration_id>', methods=['GET', 'POST'])
def payment(registration_id):
    conn = get_db()
    registration = conn.execute('SELECT * FROM registrations WHERE registration_id = ?', (registration_id,)).fetchone()
    conn.close()
    if not registration:
        return render_template('page.html', title='Registration Not Found', content='<p class="lead-dark">The registration ID could not be found.</p>'), 404
    if registration['payment_status'] == 'Paid':
        return redirect(url_for('success', registration_id=registration_id))

    if request.method == 'POST':
        reference = request.form.get('payment_reference', '').strip()
        proof = request.files.get('payment_screenshot')
        errors = []
        if not reference or len(reference) < 6:
            errors.append('Enter the UPI transaction/reference ID shown in your payment app.')
        if not proof or not proof.filename:
            errors.append('Upload a screenshot or PDF of your payment confirmation.')
        elif not allowed_proof(proof.filename):
            errors.append('Payment proof must be PNG, JPG, WEBP, or PDF.')
        if errors:
            for error in errors:
                flash(error, 'error')
            return redirect(url_for('payment', registration_id=registration_id))

        proof_name = secure_filename(proof.filename)
        stored_name = f"{registration_id}_{uuid.uuid4().hex[:8]}_{proof_name}"
        stored_path = PAYMENT_PROOF_DIR / stored_name
        proof.save(stored_path)
        submitted_at = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        conn = get_db()
        conn.execute("UPDATE registrations SET payment_reference = ?, payment_screenshot_path = ?, payment_submitted_at = ?, payment_status = 'Pending Verification', admin_note = NULL WHERE registration_id = ? AND payment_status != 'Paid'",
                     (reference, str(stored_path), submitted_at, registration_id))
        conn.commit()
        conn.close()
        flash('Payment proof submitted. Your registration is pending manual verification.', 'success')
        return redirect(url_for('success', registration_id=registration_id))

    return render_template('payment.html', registration=registration, amount=registration['amount_paise'] / 100,
                           qr_available=PAYMENT_QR_PATH.exists())


def admin_required():
    if not session.get('civilathon_admin'):
        return redirect(url_for('admin_login'))
    return None


@app.route('/admin/login', methods=['GET', 'POST'])
def admin_login():
    if request.method == 'POST':
        password = request.form.get('password', '')
        if not ADMIN_PASSWORD:
            flash('Admin password is not configured. Set CIVILATHON_ADMIN_PASSWORD in the hosting environment.', 'error')
        elif hmac_compare(password, ADMIN_PASSWORD):
            session['civilathon_admin'] = True
            return redirect(url_for('admin_dashboard'))
        else:
            flash('Incorrect password.', 'error')
    return render_template('admin_login.html')


def hmac_compare(a, b):
    import hmac
    return hmac.compare_digest(a.encode('utf-8'), b.encode('utf-8'))


@app.route('/admin/logout', methods=['POST'])
def admin_logout():
    session.pop('civilathon_admin', None)
    return redirect(url_for('admin_login'))



@app.route('/admin/clear-test-registrations', methods=['POST'])
def admin_clear_test_registrations():
    gate = admin_required()
    if gate:
        return gate
    confirmation = request.form.get('confirmation', '').strip()
    if confirmation != 'CLEAR ALL TEST REGISTRATIONS':
        flash('Registrations were not deleted. Confirmation phrase did not match.', 'error')
        return redirect(url_for('admin_dashboard'))
    conn = get_db()
    conn.execute('DELETE FROM registrations')
    conn.commit()
    conn.close()
    flash('All registration records have been cleared. Uploaded files were left untouched.', 'success')
    return redirect(url_for('admin_dashboard'))


@app.route('/admin/export-registrations.csv')
def admin_export_registrations():
    gate = admin_required()
    if gate:
        return gate
    conn = get_db()
    rows = conn.execute('''
        SELECT registration_id, leader_name, leader_mobile, leader_email,
               participation_type, team_name, college, payment_status, created_at
        FROM registrations
        ORDER BY created_at DESC
    ''').fetchall()
    conn.close()
    output = io.StringIO(newline='')
    writer = csv.writer(output)
    writer.writerow(['Registration ID', 'Participant / Team Leader', 'Mobile Number',
                     'Email', 'Participation Type', 'Team Name', 'College',
                     'Payment Status', 'Registered At'])
    for row in rows:
        writer.writerow([row['registration_id'], row['leader_name'], row['leader_mobile'],
                         row['leader_email'], row['participation_type'], row['team_name'] or '',
                         row['college'], row['payment_status'], row['created_at']])
    from flask import Response
    return Response('\ufeff' + output.getvalue(), mimetype='text/csv; charset=utf-8',
                    headers={'Content-Disposition': 'attachment; filename="civilathon_participant_contacts.csv"'})

@app.post('/admin/delete-registration/<registration_id>')
def admin_delete_registration(registration_id):
    gate = admin_required()
    if gate:
        return gate
    conn = get_db()
    row = conn.execute('SELECT * FROM registrations WHERE registration_id = ?', (registration_id,)).fetchone()
    if not row:
        conn.close()
        flash('Registration not found; nothing was deleted.', 'error')
        return redirect(url_for('admin_dashboard'))
    if row['payment_status'] == 'Paid':
        conn.close()
        flash('This registration is marked Paid and cannot be deleted here. Review it carefully first.', 'error')
        return redirect(url_for('admin_dashboard'))
    confirmation = request.form.get('confirmation', '').strip()
    if confirmation != registration_id:
        conn.close()
        flash('Registration was not deleted. Enter the exact registration ID to confirm.', 'error')
        return redirect(url_for('admin_dashboard'))
    # Delete only uploaded files belonging to this registration and inside known upload folders.
    for field, allowed_dir in [('ppt_path', UPLOAD_DIR), ('payment_screenshot_path', PAYMENT_PROOF_DIR)]:
        raw_path = row[field]
        if raw_path:
            path = Path(raw_path).resolve()
            if path.parent == allowed_dir.resolve() and path.is_file():
                try:
                    path.unlink()
                except OSError:
                    pass
    conn.execute('DELETE FROM registrations WHERE registration_id = ?', (registration_id,))
    conn.commit()
    conn.close()
    flash(f'Registration {registration_id} and its associated uploaded files were deleted.', 'success')
    return redirect(url_for('admin_dashboard'))




@app.route('/admin/export-participants.xlsx')
def admin_export_participants_xlsx():
    gate = admin_required()
    if gate:
        return gate
    conn = get_db()
    rows = conn.execute('SELECT * FROM registrations ORDER BY created_at DESC').fetchall()
    conn.close()

    # Export every registration field collected by the form and payment workflow.
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from flask import make_response

    columns = [
        ('Registration ID', 'registration_id'),
        ('Participation Type', 'participation_type'),
        ('Team Name', 'team_name'),
        ('College', 'college'),
        ('Team Leader / Participant', 'leader_name'),
        ('Leader Enrollment Number', 'leader_enrollment'),
        ('Leader Mobile', 'leader_mobile'),
        ('Leader Email', 'leader_email'),
        ('Member 2', 'member_2'),
        ('Member 2 Enrollment', 'member_2_enrollment'),
        ('Member 3', 'member_3'),
        ('Member 3 Enrollment', 'member_3_enrollment'),
        ('Member 4', 'member_4'),
        ('Member 4 Enrollment', 'member_4_enrollment'),
        ('Research Domain', 'research_domain'),
        ('Problem Statement', 'problem_statement'),
        ('Proposed Solution', 'proposed_solution'),
        ('Innovation', 'innovation'),
        ('Expected Impact', 'expected_impact'),
        ('PPT Filename', 'ppt_filename'),
        ('PPT File Path', 'ppt_path'),
        ('Payment Status', 'payment_status'),
        ('Payment ID', 'payment_id'),
        ('Payment Order ID', 'payment_order_id'),
        ('Payment Signature', 'payment_signature'),
        ('Amount (INR)', 'amount_paise'),
        ('Registered At', 'created_at'),
        ('Paid At', 'paid_at'),
        ('Payment Reference', 'payment_reference'),
        ('Payment Screenshot Path', 'payment_screenshot_path'),
        ('Payment Proof Submitted At', 'payment_submitted_at'),
        ('Admin Note', 'admin_note'),
    ]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Participants'
    sheet.append([label for label, _field in columns])
    for row in rows:
        values = []
        for label, field in columns:
            value = row[field]
            if field == 'amount_paise':
                value = (value or 0) / 100
            if value is None:
                value = ''
            values.append(value)
        sheet.append(values)

    header_fill = PatternFill('solid', fgColor='7B1F24')
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = header_fill
        cell.alignment = Alignment(vertical='top', wrap_text=True)
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for column_cells in sheet.columns:
        max_len = max(len(str(cell.value or '')) for cell in column_cells[:100])
        sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = min(max(max_len + 2, 14), 36)
    for row_cells in sheet.iter_rows(min_row=2):
        for cell in row_cells:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    buffer = io.BytesIO()
    workbook.save(buffer)
    response = make_response(buffer.getvalue())
    response.headers['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    response.headers['Content-Disposition'] = 'attachment; filename="CIVILATHON_2026_All_Participant_Details.xlsx' + '"'
    return response

@app.route('/admin/attendance-sheet.pdf')
def admin_attendance_sheet():
    gate = admin_required()
    if gate:
        return gate
    conn = get_db()
    regs = conn.execute("SELECT registration_id, participation_type, team_name, college, leader_name, leader_mobile, member_2, member_3, member_4 FROM registrations ORDER BY team_name COLLATE NOCASE, leader_name COLLATE NOCASE").fetchall()
    conn.close()
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from flask import make_response
    buffer = io.BytesIO()
    pdf = SimpleDocTemplate(buffer, pagesize=landscape(A4), rightMargin=10*mm, leftMargin=10*mm, topMargin=10*mm, bottomMargin=10*mm, title="CIVILATHON 2026 Attendance Sheet")
    styles = getSampleStyleSheet()
    title = ParagraphStyle("AttTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=15, leading=18, alignment=TA_CENTER, textColor=colors.HexColor("#7B1F24"), spaceAfter=3*mm)
    cell = ParagraphStyle("AttCell", parent=styles["Normal"], fontName="Helvetica", fontSize=8, leading=10)
    header = ParagraphStyle("AttHeader", parent=cell, fontName="Helvetica-Bold", textColor=colors.white)
    story = [Paragraph("CIVILATHON 2026 — PARTICIPANT ATTENDANCE SHEET", title), Paragraph("Civil Impact Club (CIC), LNCT Bhopal · 29 October 2026 · S-21, LNCTS New", styles["Normal"]), Spacer(1, 4*mm)]
    rows = [[Paragraph("S. No.", header), Paragraph("Team / Reg. ID", header), Paragraph("Participant Name", header), Paragraph("Mobile Number", header), Paragraph("College", header), Paragraph("Signature", header)]]
    serial = 1
    for reg in regs:
        team = reg["team_name"] or ("Individual" if reg["participation_type"] == "individual" else "—")
        people = [(reg["leader_name"], reg["leader_mobile"] or "")]
        if reg["participation_type"] == "team":
            for field in ("member_2", "member_3", "member_4"):
                name = reg[field]
                if name and name.strip():
                    people.append((name.strip(), ""))
        for name, mobile in people:
            rows.append([str(serial), Paragraph(team + "<br/><font size='7'>" + reg["registration_id"] + "</font>", cell), Paragraph(name, cell), Paragraph(mobile or " ", cell), Paragraph(reg["college"] or "", cell), ""])
            serial += 1
    if len(rows) == 1:
        rows.append(["", "No registrations found", "", "", "", ""])
    table = Table(rows, colWidths=[12*mm, 38*mm, 46*mm, 30*mm, 55*mm, 42*mm], repeatRows=1, rowHeights=[10*mm] + [13*mm]*(len(rows)-1))
    table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#7B1F24")),("GRID",(0,0),(-1,-1),0.5,colors.HexColor("#9CA3AF")),("VALIGN",(0,0),(-1,-1),"MIDDLE"),("LEFTPADDING",(0,0),(-1,-1),5),("RIGHTPADDING",(0,0),(-1,-1),5),("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#F8F5F2")])]))
    story.append(table)
    story.append(Spacer(1, 3*mm))
    story.append(Paragraph("Additional team-member mobile numbers are blank because the registration form does not collect them.", styles["Italic"]))
    pdf.build(story)
    response = make_response(buffer.getvalue())
    response.headers["Content-Type"] = "application/pdf"
    response.headers["Content-Disposition"] = 'attachment; filename="CIVILATHON_2026_Attendance_Sheet.pdf"'
    return response


@app.route('/admin')
def admin_dashboard():
    gate = admin_required()
    if gate:
        return gate
    conn = get_db()
    rows = conn.execute('SELECT * FROM registrations ORDER BY created_at DESC').fetchall()
    conn.close()
    return render_template('admin.html', registrations=rows)


@app.route('/admin/payment-proof/<registration_id>')
def admin_payment_proof(registration_id):
    gate = admin_required()
    if gate:
        return gate
    conn = get_db()
    row = conn.execute('SELECT payment_screenshot_path FROM registrations WHERE registration_id = ?', (registration_id,)).fetchone()
    conn.close()
    if not row or not row['payment_screenshot_path']:
        abort(404)
    path = Path(row['payment_screenshot_path']).resolve()
    if path.parent != PAYMENT_PROOF_DIR.resolve() or not path.is_file():
        abort(404)
    return send_from_directory(PAYMENT_PROOF_DIR, path.name, as_attachment=False)


@app.post('/admin/review/<registration_id>')
def admin_review(registration_id):
    gate = admin_required()
    if gate:
        return gate
    action = request.form.get('action', '')
    note = request.form.get('admin_note', '').strip()[:500]
    if action not in {'approve', 'reject'}:
        abort(400)
    conn = get_db()
    row = conn.execute('SELECT * FROM registrations WHERE registration_id = ?', (registration_id,)).fetchone()
    if not row:
        conn.close()
        abort(404)
    if action == 'approve':
        if not row['payment_reference'] or not row['payment_screenshot_path']:
            conn.close()
            flash('Cannot approve: payment reference and proof must be submitted first.', 'error')
            return redirect(url_for('admin_dashboard'))
        conn.execute("UPDATE registrations SET payment_status = 'Paid', paid_at = ?, admin_note = ? WHERE registration_id = ?",
                     (datetime.now().strftime('%Y-%m-%d %H:%M:%S'), note or 'Manually verified by admin', registration_id))
        flash(f'{registration_id} marked as paid. Confirm the amount in your UPI app/bank before approving.', 'success')
    else:
        conn.execute("UPDATE registrations SET payment_status = 'Rejected', admin_note = ?, paid_at = NULL WHERE registration_id = ?",
                     (note or 'Payment proof rejected. Please contact the organisers.', registration_id))
        flash(f'Payment proof for {registration_id} rejected.', 'success')
    conn.commit()
    conn.close()
    return redirect(url_for('admin_dashboard'))


@app.route('/success/<registration_id>')
def success(registration_id):
    conn = get_db()
    registration = conn.execute(
        'SELECT * FROM registrations WHERE registration_id = ?',
        (registration_id,)
    ).fetchone()
    conn.close()
    if not registration:
        return render_template('page.html', title='Registration Not Found', content='<p class="lead-dark">The registration ID could not be found.</p>'), 404
    return render_template('success.html', registration=registration)



@app.route('/submit-ppt/<registration_id>', methods=['GET', 'POST'])
def submit_ppt(registration_id):
    conn = get_db()
    registration = conn.execute(
        'SELECT * FROM registrations WHERE registration_id = ?', (registration_id,)
    ).fetchone()
    conn.close()
    if not registration:
        return render_template('page.html', title='Registration Not Found',
                              content='<p class="lead-dark">The registration ID could not be found.</p>'), 404

    # PPTs must be submitted before the event day begins.
    if datetime.now(ZoneInfo('Asia/Kolkata')).date() >= datetime(2026, 10, 29).date():
        return render_template('submit_ppt.html', registration=registration, closed=True), 410

    if request.method == 'POST':
        ppt = request.files.get('ppt')
        if not ppt or not ppt.filename:
            flash('Please choose your final PPT/PPTX file.', 'error')
            return redirect(url_for('submit_ppt', registration_id=registration_id))
        if not allowed_file(ppt.filename):
            flash('Only .ppt and .pptx files are accepted.', 'error')
            return redirect(url_for('submit_ppt', registration_id=registration_id))

        original_name = secure_filename(ppt.filename)
        stored_name = f'{registration_id}_{uuid.uuid4().hex[:8]}_{original_name}'
        stored_path = UPLOAD_DIR / stored_name
        ppt.save(stored_path)

        conn = get_db()
        conn.execute('UPDATE registrations SET ppt_filename = ?, ppt_path = ? WHERE registration_id = ?',
                     (original_name, str(stored_path), registration_id))
        conn.commit()
        conn.close()
        flash('Your PPT has been uploaded successfully.', 'success')
        return redirect(url_for('success', registration_id=registration_id))

    return render_template('submit_ppt.html', registration=registration, closed=False)


@app.route('/uploads/ppt/<path:filename>')
def uploaded_ppt(filename):
    # PPT submissions contain participant work and should be accessible only to admins.
    gate = admin_required()
    if gate:
        return gate
    safe_name = Path(filename).name
    return send_from_directory(UPLOAD_DIR, safe_name, as_attachment=True)


@app.errorhandler(413)
def too_large(_error):
    flash('PPT file is too large. Maximum allowed size is 25 MB.', 'error')
    return redirect(url_for('register'))


init_db()

if __name__ == '__main__':
    app.run(debug=os.getenv('FLASK_DEBUG', '0') == '1')
